import csv
from datetime import datetime
from enum import Enum
import io
import json
import logging
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.tenant_session import PublicAsyncSessionLocal
from app.models.reports.report_audit import ReportAudit

logger = logging.getLogger(__name__)


class BaseReportService:
    def __init__(self, db: AsyncSession, user_id: UUID, tenant_id: str):
        self.db = db
        self.user_id = user_id
        self.tenant_id = tenant_id

    async def log_export_request(
        self, report_type: str, filters: dict[str, Any], export_format: str, is_background: bool = False
    ) -> ReportAudit:
        """Log export request for audit purposes"""
        audit_record = ReportAudit(
            user_id=self.user_id,
            tenant_id=self.tenant_id,
            report_type=report_type,
            filters_applied=json.dumps(filters) if filters else None,
            export_format=export_format,
            status="pending",
            is_background_job=is_background,
        )

        # Use public database session for audit operations
        async with PublicAsyncSessionLocal() as public_db:
            public_db.add(audit_record)
            await public_db.flush()
            await public_db.commit()

        return audit_record

    async def update_audit_status(
        self,
        audit_id: UUID,
        status: str,
        file_path: str | None = None,
        file_size: int | None = None,
        error_message: str | None = None,
    ):
        """Update audit record status"""
        # Use public database session for audit operations
        async with PublicAsyncSessionLocal() as public_db:
            result = await public_db.execute(select(ReportAudit).where(ReportAudit.id == audit_id))
            audit_record = result.scalar_one_or_none()

            if audit_record:
                audit_record.status = status
                if file_path:
                    audit_record.file_path = file_path
                if file_size:
                    audit_record.file_size = file_size
                if error_message:
                    audit_record.error_message = error_message
                if status == "completed":
                    audit_record.completed_at = datetime.utcnow()
                audit_record.updated_at = datetime.utcnow()
                await public_db.commit()

    async def generate_csv_export(self, data: list[dict[str, Any]], filename: str) -> tuple[bytes, str]:
        """Generate CSV export from data"""
        if not data:
            return b"", filename

        output = io.StringIO()
        fieldnames = data[0].keys()
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(data)

        csv_content = output.getvalue().encode("utf-8")
        return csv_content, f"{filename}.csv"

    async def generate_excel_export(
        self, data: list[dict[str, Any]], filename: str, sheet_name: str = "Report"
    ) -> tuple[bytes, str]:
        """Generate Excel (.xlsx) export from data with formatting"""
        if not data:
            return b"", filename

        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Define styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        # Write headers
        if data:
            headers = list(data[0].keys())
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            # Write data rows
            for row_idx, row_data in enumerate(data, 2):
                for col_idx, header in enumerate(headers, 1):
                    value = row_data.get(header, "")
                    # Handle None values and format dates
                    if value is None:
                        value = ""
                    elif isinstance(value, datetime):
                        value = value.strftime("%Y-%m-%d %H:%M:%S")
                    elif isinstance(value, Enum):
                        value = value.value
                    elif isinstance(value, UUID):
                        value = str(value)
                    ws.cell(row=row_idx, column=col_idx, value=value)

            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
                ws.column_dimensions[column_letter].width = adjusted_width

        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        excel_content = output.getvalue()
        output.close()

        return excel_content, f"{filename}.xlsx"

    async def generate_pdf_export(
        self, data: list[dict[str, Any]], filename: str, title: str = "Report", subtitle: str = None
    ) -> tuple[bytes, str]:
        """Generate PDF export from data with professional formatting"""
        if not data:
            return b"", filename

        # Create PDF in memory
        pdf_buffer = io.BytesIO()
        doc = SimpleDocTemplate(pdf_buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)

        # Define styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Heading1"],
            fontSize=16,
            spaceAfter=30,
            alignment=TA_CENTER,
            textColor=colors.darkblue,
        )

        subtitle_style = ParagraphStyle(
            "CustomSubtitle",
            parent=styles["Heading2"],
            fontSize=12,
            spaceAfter=20,
            alignment=TA_CENTER,
            textColor=colors.grey,
        )

        # Build PDF content
        story = []

        # Add title
        story.append(Paragraph(title, title_style))

        # Add subtitle if provided
        if subtitle:
            story.append(Paragraph(subtitle, subtitle_style))

        # Add generation timestamp
        timestamp = datetime.now().strftime("%B %d, %Y at %I:%M %p")
        story.append(Paragraph(f"Generated on: {timestamp}", styles["Normal"]))
        story.append(Spacer(1, 20))

        # Prepare table data
        if data:
            headers = list(data[0].keys())

            # Create table data with headers
            table_data = [headers]

            # Add data rows
            for row in data:
                row_data = []
                for header in headers:
                    value = row.get(header, "")
                    # Handle None values and format dates
                    if value is None:
                        value = ""
                    elif isinstance(value, datetime):
                        value = value.strftime("%Y-%m-%d %H:%M:%S")
                    else:
                        value = str(value)
                    row_data.append(value)
                table_data.append(row_data)

            # Create table
            table = Table(table_data, repeatRows=1)

            # Apply table style
            table_style = TableStyle(
                [
                    # Header row styling
                    ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 10),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                    # Data rows styling
                    ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 1), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )

            table.setStyle(table_style)
            story.append(table)

            # Add record count
            story.append(Spacer(1, 20))
            story.append(Paragraph(f"Total Records: {len(data)}", styles["Normal"]))

        # Build PDF
        doc.build(story)

        # Get PDF content
        pdf_content = pdf_buffer.getvalue()
        pdf_buffer.close()

        return pdf_content, f"{filename}.pdf"

    async def create_background_export_job(
        self,
        report_type: str,
        filters: dict[str, Any],
        export_format: str,
        filename: str,
        user_id: UUID,
        tenant_id: str,
    ) -> tuple[str, str]:
        """
        Create a background job for export and return job ID and audit ID
        """
        # Create audit record
        audit_id = await self.log_export_request(
            report_type=report_type,
            filters=filters,
            export_format=export_format,
            filename=filename,
            user_id=user_id,
            tenant_id=tenant_id,
            is_background=True,
        )

        # Start background task
        from app.tasks.report_tasks import generate_staff_report_export, generate_student_report_export

        if report_type.startswith("student"):
            task = generate_student_report_export.delay(
                audit_id=str(audit_id),
                user_id=str(user_id),
                tenant_id=tenant_id,
                report_type=report_type,
                filters=filters,
                export_format=export_format,
                filename=filename,
            )
        elif report_type.startswith("staff"):
            task = generate_staff_report_export.delay(
                audit_id=str(audit_id),
                user_id=str(user_id),
                tenant_id=tenant_id,
                report_type=report_type,
                filters=filters,
                export_format=export_format,
                filename=filename,
            )
        else:
            raise ValueError(f"Unsupported report type for background export: {report_type}")

        return str(task.id), str(audit_id)

    async def fetch_all_rows(self, fetch, filters) -> tuple[list[dict[str, Any]], int]:
        """Run a paginated report fetch and return every matching row, ignoring the filter's page and page_size"""
        data, total_count = await fetch(filters)
        if total_count > len(data):
            data, total_count = await fetch(filters.model_copy(update={"page": 1, "page_size": total_count}))
        return data, total_count

    def should_use_background_job(self, data_count: int, export_format: str) -> bool:
        """
        Determine if export should use background job based on data size and format
        """
        # Use background job for:
        # - Large datasets (>1000 records)
        # - Excel exports with >500 records
        # - PDF exports with >300 records
        if data_count > 1000:
            return True
        elif export_format == "xlsx" and data_count > 500:
            return True
        elif export_format == "pdf" and data_count > 300:
            return True

        return False

    def apply_pagination(self, query, page: int = 1, page_size: int = 100) -> tuple[Any, int]:
        """Apply pagination to query"""
        offset = (page - 1) * page_size
        paginated_query = query.offset(offset).limit(page_size)
        return paginated_query, page_size

    def apply_filters(self, query, filters: dict[str, Any], filter_mappings: dict[str, Any]):
        """Apply filters to query based on mappings"""
        for filter_key, filter_value in filters.items():
            if filter_value is not None and filter_key in filter_mappings:
                column = filter_mappings[filter_key]
                if isinstance(filter_value, list):
                    query = query.where(column.in_(filter_value))
                elif isinstance(filter_value, dict) and "min" in filter_value and "max" in filter_value:
                    query = query.where(column >= filter_value["min"], column <= filter_value["max"])
                else:
                    query = query.where(column == filter_value)
        return query

    def apply_sorting(self, query, sort_by: str | None = None, sort_order: str = "asc"):
        """Apply sorting to query"""
        if sort_by:
            if hasattr(query.column_descriptions[0]["entity"], sort_by):
                column = getattr(query.column_descriptions[0]["entity"], sort_by)
                if sort_order.lower() == "desc":
                    query = query.order_by(column.desc())
                else:
                    query = query.order_by(column.asc())
        return query

    async def validate_export_limits(
        self, row_count: int, max_rows: int = 5000, max_file_size: int = 50 * 1024 * 1024  # 50MB
    ) -> tuple[bool, str]:
        """Validate export limits"""
        if row_count > max_rows:
            return (
                False,
                f"Export exceeds maximum row limit of {max_rows:,} rows. Use background export for larger datasets.",
            )

        estimated_size = row_count * 100  # Rough estimate: 100 bytes per row
        if estimated_size > max_file_size:
            return False, f"Estimated file size exceeds maximum limit of {max_file_size // (1024*1024)}MB"

        return True, ""

    def get_tenant_id(self) -> str:
        """Get current tenant id"""
        return self.tenant_id

    async def execute_tenant_query(self, query_text: str, params: dict[str, Any] = None):
        """Execute query in the tenant session (row-level security scopes it to the tenant)"""
        if params:
            result = await self.db.execute(text(query_text), params)
        else:
            result = await self.db.execute(text(query_text))

        return result
