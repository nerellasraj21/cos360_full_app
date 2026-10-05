import logging
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID

import openpyxl
from fastapi import HTTPException
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.auth.role_model import Role
from app.models.auth.user_model import User
from app.models.masters.designations_model import Designation
from app.models.masters.staff_model import QualificationLevelEnum, Staff
from app.schemas.masters.staff_schema import StaffEnrollmentCreate, StaffQualificationCreate
from app.service.masters.staff_service import add_staff_qualification, create_staff_enrollment

logger = logging.getLogger(__name__)

SHEET_NAME = "Staff Admission"

TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2] / "static" / "templates" / "staff_bulk_upload_template.xlsx"
)

# Mandatory fields for staff bulk upload: First Name, Phone, Address
REQUIRED_HEADERS = [
    "First Name",
    "Phone",
    "Address",
]

DATE_FORMATS = ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y")

QUALIFICATION_LEVEL_ALIASES = {
    "below graduation": QualificationLevelEnum.below_graduation,
    "below_graduation": QualificationLevelEnum.below_graduation,
    "graduation": QualificationLevelEnum.graduation,
    "post graduation": QualificationLevelEnum.post_graduation,
    "post-graduation": QualificationLevelEnum.post_graduation,
    "post_graduation": QualificationLevelEnum.post_graduation,
    "phd": QualificationLevelEnum.phd,
}


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_int(value: Any) -> int | None:
    text = _clean(value)
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _parse_decimal(value: Any) -> Decimal | None:
    text = _clean(value)
    if not text:
        return None
    try:
        return Decimal(text.replace(",", ""))
    except InvalidOperation:
        return None


def _map_account_type(value: Any) -> str | None:
    text = _clean(value)
    if not text:
        return None
    normalized = text.capitalize()
    return normalized if normalized in {"Savings", "Current"} else None


class _RowError(Exception):
    pass


def _map_qualification_level(value: str) -> QualificationLevelEnum:
    level = QUALIFICATION_LEVEL_ALIASES.get(value.lower())
    if not level:
        valid = ", ".join(e.value for e in QualificationLevelEnum)
        raise _RowError(f"Invalid Qualification Level '{value}'. Must be one of: {valid}")
    return level


async def _get_designation_id(
    db: AsyncSession, title: str | None, cache: dict[str, UUID]
) -> UUID | None:
    if not title:
        return None
    key = title.lower()
    if key in cache:
        return cache[key]
    result = await db.execute(select(Designation).where(func.lower(Designation.title) == key))
    designation = result.scalars().first()
    if not designation:
        raise _RowError(f"Designation '{title}' not found")
    cache[key] = designation.id
    return designation.id


async def _get_role_id(db: AsyncSession, name: str | None, cache: dict[str, UUID]) -> UUID | None:
    if not name:
        return None
    key = name.lower()
    if key in cache:
        return cache[key]
    result = await db.execute(select(Role).where(func.lower(Role.name) == key))
    role = result.scalars().first()
    if not role:
        raise _RowError(f"Role '{name}' not found")
    cache[key] = role.id
    return role.id


async def _check_duplicate_email(db: AsyncSession, email: str) -> None:
    result = await db.execute(
        select(User.id).where(or_(func.lower(User.email) == email.lower(), User.username == email))
    )
    if result.first():
        raise _RowError(f"A user with email '{email}' already exists")
    result = await db.execute(select(Staff.id).where(func.lower(Staff.email) == email.lower()))
    if result.first():
        raise _RowError(f"A staff member with email '{email}' already exists")


def _build_row_lookup(header_row: list[Any]) -> dict[str, int]:
    lookup: dict[str, int] = {}
    for idx, value in enumerate(header_row):
        if value is None:
            continue
        header = str(value).strip()
        if header:
            lookup[header] = idx
    return lookup


async def parse_and_bulk_create_staff_enrollments(file_bytes: bytes, db: AsyncSession) -> dict:
    """
    Parse the "Staff Admission" sheet of the staff bulk-upload Excel template and
    create one staff enrollment per row via the existing `create_staff_enrollment`
    flow, so all normal business rules (user account creation, default 'Staff'
    role lookup, gender validation, etc.) still apply.

    Mandatory columns: First Name, Phone, Address.
    Email, Joining Date, Qualification Level and Degree/Course are optional.

    Returns {"created": [...], "errors": [...], "total_rows": int}
    """
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception as e:
        logger.error("Unhandled error: %s", e)
        raise HTTPException(400, detail="Invalid Excel file")

    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    header_row = [cell.value for cell in ws[1]]
    col_index = _build_row_lookup(header_row)

    missing = [h for h in REQUIRED_HEADERS if h not in col_index]
    if missing:
        raise HTTPException(400, detail=f"Missing required column(s): {', '.join(missing)}")

    def cell(row: tuple, name: str) -> Any:
        idx = col_index.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    designation_cache: dict[str, UUID] = {}
    role_cache: dict[str, UUID] = {}

    created: list[dict] = []
    errors: list[str] = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        try:
            first_name = _clean(cell(row, "First Name"))
            if not first_name:
                raise _RowError("First Name is required")

            email = _clean(cell(row, "Email"))

            phone = _clean(cell(row, "Phone"))
            if not phone:
                raise _RowError("Phone is required")

            address = _clean(cell(row, "Address"))
            if not address:
                raise _RowError("Address is required")

            if email:
                await _check_duplicate_email(db, email)

            designation_id = await _get_designation_id(
                db, _clean(cell(row, "Designation")), designation_cache
            )
            role_id = await _get_role_id(db, _clean(cell(row, "Role")), role_cache)

            # Optional qualification row: level and degree must come together
            level_text = _clean(cell(row, "Qualification Level"))
            degree_text = _clean(cell(row, "Degree/Course"))
            qualification_data = None
            if level_text or degree_text:
                if not (level_text and degree_text):
                    raise _RowError("Qualification Level and Degree/Course must be provided together")
                qualification_data = StaffQualificationCreate(
                    level=_map_qualification_level(level_text),
                    name=degree_text,
                    passed_out_year=_parse_int(cell(row, "Pass-out year")),
                    percentage=_parse_decimal(cell(row, "Percentage/CGPA")),
                    university=_clean(cell(row, "University/Board")),
                )

            enrollment = StaffEnrollmentCreate(
                first_name=first_name,
                last_name=_clean(cell(row, "Last Name")),
                email=email,
                phone=phone,
                gender=_clean(cell(row, "Gender")),
                date_of_birth=_parse_date(cell(row, "Date of birth")),
                joining_date=_parse_date(cell(row, "Joining Date")),
                qualification=_clean(cell(row, "Qualification")),
                experience_years=_parse_int(cell(row, "Experience")),
                address=address,
                designation_id=designation_id,
                department=_clean(cell(row, "Department")),
                role_id=role_id,
                # Work Experience
                work_org=_clean(cell(row, "Previous organization")),
                subjects_dealt=_clean(cell(row, "Subjects dealt")),
                work_from_date=_parse_date(cell(row, "From date")),
                work_to_date=_parse_date(cell(row, "To date")),
                work_remarks=_clean(cell(row, "Remarks")),
                # Bank Details
                bank_name=_clean(cell(row, "Bank Name")),
                bank_branch=_clean(cell(row, "Branch")),
                account_number=_clean(cell(row, "Account number")),
                ifsc_code=_clean(cell(row, "IFSC Code")),
                account_holder_name=_clean(cell(row, "Account Holder Name")),
                account_type=_map_account_type(cell(row, "Account type")),
                # Salary & PF
                last_drawn_salary=_parse_decimal(cell(row, "Last Drawn Salary")),
                pf_account_number=_clean(cell(row, "PF Account Number")),
                uan_number=_clean(cell(row, "UAN Number")),
            )

            staff = await create_staff_enrollment(enrollment, db)

            if qualification_data:
                await add_staff_qualification(staff.id, qualification_data, db)

            created.append(
                {
                    "row": row_idx,
                    "staff_id": str(staff.id),
                    "name": f"{first_name} {enrollment.last_name or ''}".strip(),
                    "email": email,
                }
            )
        except _RowError as e:
            errors.append(f"Row {row_idx}: {str(e)}")
        except ValidationError as e:
            first_error = e.errors()[0]
            field = ".".join(str(loc) for loc in first_error.get("loc", []))
            errors.append(f"Row {row_idx}: {field} - {first_error.get('msg')}")
        except HTTPException as e:
            detail = e.detail
            if isinstance(detail, dict):
                detail = detail.get("message") or detail.get("detail") or str(detail)
            errors.append(f"Row {row_idx}: {detail}")
        except Exception as e:
            logger.error(f"Unexpected error on row {row_idx} of staff bulk upload: {str(e)}")
            errors.append(f"Row {row_idx}: {str(e)}")

    return {"created": created, "errors": errors, "total_rows": len(created) + len(errors)}


def _add_dropdown_list(
    ws, col_index: dict[str, int], list_ws, list_col_idx: int, values: list[str], target_headers: list[str]
) -> None:
    """Write `values` into a hidden helper sheet column and wire a list-type
    DataValidation on `target_headers` referencing that column, so the
    dropdown always reflects live master data instead of a fixed guess."""
    if not values:
        return
    for row_idx, value in enumerate(values, start=1):
        list_ws.cell(row=row_idx, column=list_col_idx, value=value)
    list_letter = get_column_letter(list_col_idx)
    dv = DataValidation(
        type="list",
        formula1=f"'{list_ws.title}'!${list_letter}$1:${list_letter}${len(values)}",
        allow_blank=True,
        showErrorMessage=True,
        showInputMessage=True,
        errorStyle="stop",
        errorTitle="Invalid value",
        error="Choose a value from the dropdown list.",
    )
    ws.add_data_validation(dv)
    for header in target_headers:
        idx = col_index.get(header)
        if idx is None:
            continue
        letter = get_column_letter(idx + 1)
        dv.add(f"{letter}2:{letter}1000")


async def generate_blank_staff_template(db: AsyncSession) -> bytes:
    """
    Load the blank staff bulk-upload template with live Designation/Role
    dropdowns wired in, so uploaders can only pick names that currently exist
    for this tenant instead of mistyping one that will fail on upload.
    """
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    header_row = [cell.value for cell in ws[1]]
    col_index = _build_row_lookup(header_row)

    designation_titles = [row[0] for row in (await db.execute(select(Designation.title).order_by(Designation.title))).all()]
    role_names = [row[0] for row in (await db.execute(select(Role.name).order_by(Role.name))).all()]

    list_ws = wb["Lists"] if "Lists" in wb.sheetnames else wb.create_sheet("Lists")
    list_ws.sheet_state = "hidden"

    _add_dropdown_list(ws, col_index, list_ws, 1, designation_titles, ["Designation"])
    _add_dropdown_list(ws, col_index, list_ws, 2, role_names, ["Role"])

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
