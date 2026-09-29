from io import BytesIO
from uuid import UUID

from fastapi import HTTPException
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.exam.exam_subject_config_model import ExamSubjectConfig


async def generate_excel_template(
    exam_id: UUID,
    subject_config_id: UUID,
    db: AsyncSession,
    class_id: UUID = None,
    section_id: UUID = None,
    attempt_number: int = 1,
) -> bytes:
    """
    Generate an Excel template for offline mark entry.
    Columns: student_id, Roll No, Student Name, one column per component, Remarks.
    Pre-populates enrolled student rows when class_id is provided.
    Returns bytes of the xlsx file.
    """
    from sqlalchemy import text

    # Get subject config and components
    result = await db.execute(
        select(ExamSubjectConfig)
        .options(selectinload(ExamSubjectConfig.components))
        .where(ExamSubjectConfig.id == subject_config_id)
    )
    config = result.scalar_one_or_none()
    if not config:
        raise HTTPException(404, detail="Subject config not found")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Marks Entry"

    # Build header row
    headers = ["student_id", "Roll No", "Student Name"]
    for comp in config.components:
        headers.append(f"{comp.component_name} (Max: {comp.max_marks})")
    headers.append("Remarks")

    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Freeze the header row so it stays visible when scrolling
    ws.freeze_panes = "A2"

    # Pre-populate student rows when class_id is provided
    if class_id:
        section_filter = "AND sa.current_section_id = :section_id" if section_id else ""
        student_sql = text(f"""
            SELECT student_id, admission_number, student_name
            FROM (
                SELECT DISTINCT ON (sa.student_id)
                    sa.student_id,
                    sa.admission_number,
                    TRIM(COALESCE(s.first_name, '') || ' ' || COALESCE(s.last_name, '')) AS student_name
                FROM student_admissions sa
                JOIN students s ON s.id = sa.student_id
                WHERE sa.current_class_id = :class_id
                  {section_filter}
                ORDER BY sa.student_id, sa.admission_number
            ) sub
            ORDER BY admission_number
        """)
        params: dict = {"class_id": str(class_id)}
        if section_id:
            params["section_id"] = str(section_id)
        student_rows = (await db.execute(student_sql, params)).fetchall()

        # Fetch already-entered marks for this exam + subject config
        from sqlalchemy import and_

        from app.models.exam.student_marks_model import StudentMark

        student_ids = [r[0] for r in student_rows]
        mark_rows_result = await db.execute(
            StudentMark.__table__.select().where(
                and_(
                    StudentMark.exam_id == exam_id,
                    StudentMark.subject_config_id == subject_config_id,
                    StudentMark.attempt_number == attempt_number,
                    StudentMark.student_id.in_(student_ids),
                )
            )
        )
        # Build lookup: student_id → component_id → {"marks": value, "is_absent": bool}
        mark_lookup: dict = {}
        for mr in mark_rows_result.fetchall():
            sid_key = str(mr.student_id)
            cid_key = str(mr.component_id)
            if sid_key not in mark_lookup:
                mark_lookup[sid_key] = {}
            mark_lookup[sid_key][cid_key] = {
                "marks": mr.marks_obtained,
                "is_absent": bool(mr.is_absent),
            }

        data_font = Font(name="Calibri", size=11)
        id_font = Font(name="Calibri", size=11, color="000000")

        for row_idx, (sid, adm_no, name) in enumerate(student_rows, start=2):
            ws.cell(row=row_idx, column=1, value=str(sid)).font = id_font
            ws.cell(row=row_idx, column=2, value=adm_no or "").font = data_font
            ws.cell(row=row_idx, column=3, value=name or "").font = data_font
            # Fill in any marks already entered for this student
            student_marks = mark_lookup.get(str(sid), {})
            for comp_col_offset, comp in enumerate(config.components):
                entry = student_marks.get(str(comp.id))
                col_idx = 4 + comp_col_offset
                if entry is not None:
                    if entry["is_absent"]:
                        cell = ws.cell(row=row_idx, column=col_idx, value="ABS")
                        cell.font = data_font
                        cell.alignment = Alignment(horizontal="right")
                    elif entry["marks"] is not None:
                        ws.cell(row=row_idx, column=col_idx, value=float(entry["marks"])).font = data_font

    # Auto-size columns (approximate — openpyxl has no auto-fit)
    for col_idx, header in enumerate(headers, 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(len(str(header)) + 4, 14)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


async def parse_excel_upload(
    file_bytes: bytes,
    exam_id: UUID,
    subject_config_id: UUID,
    entered_by: UUID,
    db: AsyncSession,
) -> dict:
    """
    Parse uploaded Excel file and return mark entry payload.
    Returns {"marks": [...], "errors": [...], "total_rows": int}

    Expected sheet layout (produced by generate_excel_template):
      Row 1 — headers: student_id | Roll No | Student Name | <component cols> | Remarks
      Row 2+ — one student per row.

    The "student_id" column (index 0) is the authoritative identifier.
    Component mark columns start at index 3 (after student_id, Roll No, Student Name).
    The final "Remarks" column is captured but not mapped to components.
    """
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes))
        ws = wb.active
    except Exception as e:
        raise HTTPException(400, detail=f"Invalid Excel file: {str(e)}")

    marks = []
    errors = []
    headers = [cell.value for cell in ws[1]]

    # Derive component column names (everything between index 3 and the last "Remarks" col)
    component_headers = headers[3:-1] if headers and headers[-1] == "Remarks" else headers[3:]

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        try:
            student_id = row[0]
            if not student_id:
                errors.append(f"Row {row_idx}: missing student_id")
                continue

            # Map component header → raw cell value for this row
            component_data = dict(zip(component_headers, row[3 : 3 + len(component_headers)], strict=False))
            remarks_value = row[3 + len(component_headers)] if len(row) > 3 + len(component_headers) else None

            marks.append(
                {
                    "student_id": str(student_id),
                    "roll_no": row[1],
                    "student_name": row[2],
                    "row_data": component_data,
                    "remarks": remarks_value,
                }
            )
        except Exception as e:
            errors.append(f"Row {row_idx}: {str(e)}")

    return {"marks": marks, "errors": errors, "total_rows": len(marks)}
